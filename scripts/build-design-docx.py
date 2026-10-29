"""Build the submission-ready Word version of the concise design note."""
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "DESIGN.md"
OUTPUT = ROOT / "docs" / "DESIGN.docx"


def set_cell_border(cell, **kwargs):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, attrs in kwargs.items():
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key, value in attrs.items():
            element.set(qn("w:" + key), str(value))


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def add_inline(paragraph, text):
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        run = paragraph.add_run(part.strip("`") if part.startswith("`") else part)
        if part.startswith("`"):
            run.font.name = "Courier New"
            run.font.size = Pt(9)


def remove_paragraph_border(paragraph):
    paragraph_properties = paragraph._p.get_or_add_pPr()
    border = paragraph_properties.find(qn("w:pBdr"))
    if border is not None:
        paragraph_properties.remove(border)


def main():
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = document.styles
    title_style = styles["Title"]
    style_border = title_style._element.pPr.find(qn("w:pBdr")) if title_style._element.pPr is not None else None
    if style_border is not None:
        title_style._element.pPr.remove(style_border)
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.3)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.03
    for name, size, before, after in [("Heading 1", 13, 8, 3), ("Heading 2", 11.2, 6, 2), ("Heading 3", 10.4, 4, 1)]:
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Heading 1" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title_run = title.add_run("NFPL Access Control Service Design Note")
    title_run.font.name = "Aptos Display"
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(2)
    remove_paragraph_border(title)
    subtitle = document.add_paragraph("Concise submission design document - implementation-aligned")
    subtitle.runs[0].italic = True
    subtitle.runs[0].font.size = Pt(9.5)
    subtitle.paragraph_format.space_after = Pt(8)

    header = section.header.paragraphs[0]
    header.text = "NFPL Access Control Service"
    header.runs[0].font.size = Pt(8.5)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    footer.runs[0].font.size = Pt(8.5)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    skip_title = True
    for line in lines:
        if line.startswith("# ") and skip_title:
            skip_title = False
            continue
        if not line.strip():
            continue
        if line.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            paragraph.add_run(line[3:])
        elif line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 3")
            paragraph.add_run(line[4:])
        else:
            paragraph = document.add_paragraph()
            add_inline(paragraph, line)

    document.core_properties.title = "NFPL Access Control Service Design Note"
    document.core_properties.subject = "Concise authorization design document"
    document.core_properties.author = "NFPL"
    document.save(OUTPUT)


if __name__ == "__main__":
    main()
